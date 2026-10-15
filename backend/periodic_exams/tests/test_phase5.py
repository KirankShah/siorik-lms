import csv
import io
from datetime import date, timedelta
from decimal import Decimal

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken
from django.test import TestCase

from accounts.models import Organization, User
from periodic_exams.models import PeriodicExamAttempt
from periodic_exams.services import ensure_current_cycle

from .test_phase2 import NOW, add_question, enable, set_joined, BEFORE_ALL_TEST_CYCLES


class PeriodicExamReportTests(TestCase):
    def setUp(self):
        self.organization = Organization.objects.create(name='Report Bank', slug='report-bank')
        settings = enable(self.organization, questions_per_exam=1)
        settings.pass_mark_percent = 60
        settings.save(update_fields=['pass_mark_percent'])
        add_question(self.organization, User.AssessmentLevel.OFFICER, 'Report question')
        self.learners = {}
        for key, first_name in (
            ('none', '=Formula'), ('progress', 'प्रगति'), ('failed', 'Failed'),
            ('passed', 'Passed'), ('deleted', 'Deleted'), ('inactive', 'Inactive'),
        ):
            learner = User.objects.create_user(
                email=f'{key}@example.com', password='Pass1234', role=User.Role.LEARNER,
                organization=self.organization, assessment_level=User.AssessmentLevel.OFFICER,
                first_name=first_name, last_name='Learner', branch_department='@Operations',
            )
            set_joined(learner, BEFORE_ALL_TEST_CYCLES)
            self.learners[key] = learner
        self.cycle = ensure_current_cycle(self.organization, NOW)
        self._attempt('progress', 1, submitted=False)
        self._attempt('failed', 1, submitted=True, score=Decimal('40.00'))
        self._attempt('passed', 1, submitted=True, score=Decimal('40.00'))
        self._attempt('passed', 2, submitted=True, score=Decimal('80.00'), passed=True)
        self.learners['deleted'].delete()
        self.learners['inactive'].is_active = False
        self.learners['inactive'].save(update_fields=['is_active'])

        self.admin = User.objects.create_user(
            email='report-admin@example.com', password='Pass1234', role=User.Role.ORG_ADMIN,
            organization=self.organization,
        )
        self.platform = User.objects.create_user(
            email='platform-report@example.com', password='Pass1234', role=User.Role.PLATFORM_ADMIN,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def _attempt(self, key, number, *, submitted, score=Decimal('0.00'), passed=False):
        participant = self.cycle.participants.get(email_snapshot=f'{key}@example.com')
        started = NOW + timedelta(minutes=number)
        PeriodicExamAttempt.objects.create(
            participant=participant, attempt_number=number, started_at=started,
            timer_segment_started_at=started, submitted_at=started if submitted else None,
            score_percent=score, passed=passed,
            submission_reason='LEARNER' if submitted else None,
        )
        if submitted:
            participant.attempts_used = number
            participant.best_score_percent = max(
                participant.best_score_percent or Decimal('0.00'), score
            )
            participant.last_attempt_at = started
            if passed:
                participant.passed_at = started
            participant.save(update_fields=[
                'attempts_used', 'best_score_percent', 'last_attempt_at', 'passed_at'
            ])

    def url(self, suffix=''):
        return f'/api/periodic-exams/report/cycles/{self.cycle.id}/{suffix}'

    def test_hand_calculated_rows_search_status_snapshots_and_empty_cycle(self):
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200, response.data)
        rows = {row['email']: row for row in response.data['results']}
        self.assertEqual(response.data['count'], 6)
        self.assertEqual(rows['none@example.com']['status'], 'Not attempted')
        self.assertFalse(rows['none@example.com']['attempted'])
        self.assertEqual(rows['progress@example.com']['status'], 'In progress')
        self.assertFalse(rows['progress@example.com']['completed'])
        self.assertEqual(rows['failed@example.com']['status'], 'Failed')
        self.assertEqual(rows['passed@example.com']['status'], 'Passed')
        self.assertEqual(rows['passed@example.com']['attempts_used'], 2)
        self.assertEqual(rows['passed@example.com']['best_score'], 80.0)
        self.assertEqual(rows['deleted@example.com']['staff_name'], 'Deleted Learner')
        self.assertIn('inactive@example.com', rows)
        self.assertEqual(rows['progress@example.com']['staff_name'], 'प्रगति Learner')
        self.assertEqual(self.client.get(self.url(), {'search': 'प्रगति'}).data['count'], 1)
        self.assertEqual(self.client.get(self.url(), {'status': 'Passed'}).data['count'], 1)

        empty_org = Organization.objects.create(name='Empty Report', slug='empty-report')
        enable(empty_org, questions_per_exam=1)
        add_question(empty_org, User.AssessmentLevel.OFFICER, 'Empty question')
        empty_cycle = ensure_current_cycle(empty_org, NOW)
        empty_admin = User.objects.create_user(
            email='empty-admin@example.com', password='Pass1234', role=User.Role.ORG_ADMIN,
            organization=empty_org,
        )
        self.client.force_authenticate(empty_admin)
        self.assertEqual(
            self.client.get(f'/api/periodic-exams/report/cycles/{empty_cycle.id}/').data['count'], 0
        )

    def test_json_csv_xlsx_parity_formula_safety_and_nepali_unicode(self):
        json_rows = self.client.get(self.url()).data['results']
        csv_response = self.client.get(f'/api/periodic-exams/report/cycles/{self.cycle.id}.csv')
        csv_text = b''.join(csv_response.streaming_content).decode('utf-8-sig')
        csv_rows = list(csv.reader(io.StringIO(csv_text)))
        self.assertEqual(csv_rows[0], [
            'Staff name', 'Email', 'Level', 'Branch/department', 'Cycle', 'Attempted',
            'Completed', 'Attempts used', 'Best score', 'Status', 'Last attempt date',
        ])
        self.assertEqual(len(csv_rows) - 1, len(json_rows))
        formula_row = next(row for row in csv_rows[1:] if row[1] == 'none@example.com')
        self.assertEqual(formula_row[0], "'=Formula Learner")
        self.assertEqual(formula_row[3], "'@Operations")
        self.assertIn('प्रगति Learner', csv_text)

        xlsx_response = self.client.get(f'/api/periodic-exams/report/cycles/{self.cycle.id}.xlsx')
        workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
        values = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(list(values[0]), csv_rows[0])
        self.assertEqual(len(values) - 1, len(json_rows))
        xlsx_formula = next(row for row in values[1:] if row[1] == 'none@example.com')
        self.assertEqual(xlsx_formula[0], "'=Formula Learner")
        self.assertIn('प्रगति Learner', [row[0] for row in values[1:]])

    def test_tenant_platform_selector_cycle_ownership_and_role_denials(self):
        other_org = Organization.objects.create(name='Other Report', slug='other-report')
        enable(other_org, questions_per_exam=1)
        add_question(other_org, User.AssessmentLevel.OFFICER, 'Other report question')
        other_cycle = ensure_current_cycle(other_org, NOW)
        self.assertEqual(self.client.get(self.url()).status_code, 200)
        self.assertEqual(
            self.client.get(f'/api/periodic-exams/report/cycles/{other_cycle.id}/').status_code, 404
        )
        self.client.force_authenticate(self.platform)
        self.assertEqual(self.client.get(self.url()).status_code, 400)
        selected = self.client.get(self.url(), {'organization_id': self.organization.id})
        self.assertEqual(selected.status_code, 200)
        self.assertEqual(selected.data['count'], 6)
        for role in (User.Role.LEARNER, User.Role.INSTRUCTOR):
            user = User.objects.create_user(
                email=f'{role.lower()}-report@example.com', password='Pass1234', role=role,
                organization=self.organization,
            )
            self.client.force_authenticate(user)
            self.assertEqual(self.client.get(self.url()).status_code, 403)

    def test_subscription_lock_blocks_org_admin_but_not_platform_admin(self):
        org_settings = self.organization.settings
        org_settings.subscription_start_date = date.today() - timedelta(days=10)
        org_settings.subscription_duration_days = 1
        org_settings.org_admin_grace_period_days = 0
        org_settings.save(update_fields=[
            'subscription_start_date', 'subscription_duration_days', 'org_admin_grace_period_days'
        ])
        self.client.force_authenticate(user=None)
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(self.admin)}')
        self.assertEqual(self.client.get(self.url()).status_code, 403)
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(self.platform)}')
        self.assertEqual(
            self.client.get(self.url(), {'organization_id': self.organization.id}).status_code, 200
        )
